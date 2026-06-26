"""Parse LandslideInventory.xlsx ('landslides' sheet) into a clean events table.

Column layout (0-indexed), confirmed against the 2024 file:
  row 0 = group headers, row 1 = column headers, data starts row 2.
   0 Name | 1 Date(UTC) | 2 Time(UTC) picked | 3 Time(UTC) grid | 4 Date(AKTZ) | 5 Time(AKTZ)
   6 GridCenter Lat | 7 GridCenter Lon | 8 GridCenter Ref
   9 GroundTruth Y/N | 10 GT Lat | 11 GT Lon | 12 GT Volume | 13 GT Ref
  14 In ESEC Y/N | 15 Nearest station | 16 Station dist (km)
  17 GridSearch Lat | 18 GridSearch Lon | 19 Coherency | 20 AbsLocErr km
  21 RelErrMin km | 22 RelErrMax km | 23 Seismic Volume m3
  28 Real-time Y/N | 29 Notes
"""
from __future__ import annotations
import datetime as dt
import math
import pandas as pd
from openpyxl import load_workbook


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(float(v)) else None
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "?", "CHECK") or s.lower().startswith(("ask", "need")):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _time(v):
    """Cell may be a datetime.time, a fraction of a day, or a string."""
    if v is None or str(v).strip() in ("", "-"):
        return None
    if isinstance(v, dt.time):
        return v
    if isinstance(v, dt.datetime):
        return v.time()
    if isinstance(v, (int, float)):
        secs = round(float(v) * 86400)
        return dt.time(secs // 3600 % 24, secs % 3600 // 60, secs % 60)
    try:
        return dt.datetime.strptime(str(v).strip(), "%H:%M:%S").time()
    except ValueError:
        return None


def _date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):  # Excel serial, 1900 system
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    return None


def load_events(path: str, sheet: str = "landslides") -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    out = []
    for r in rows:
        if not r or r[0] is None or str(r[0]).strip() == "":
            continue
        r = list(r) + [None] * (30 - len(r))
        date = _date(r[1])
        t = _time(r[2]) or _time(r[3])  # prefer picked time, fall back to grid time
        when = dt.datetime.combine(date, t or dt.time(0, 0)) if date else None

        gt_lat, gt_lon = _num(r[10]), _num(r[11])
        gc_lat, gc_lon = _num(r[6]), _num(r[7])
        gs_lat, gs_lon = _num(r[17]), _num(r[18])
        if gt_lat is not None and gt_lon is not None:
            lat, lon, src = gt_lat, gt_lon, "ground_truth"
        elif gc_lat is not None and gc_lon is not None:
            lat, lon, src = gc_lat, gc_lon, "grid_center"
        elif gs_lat is not None and gs_lon is not None:
            lat, lon, src = gs_lat, gs_lon, "grid_search"
        else:
            lat = lon = None
            src = "none"

        abs_err = _num(r[20]) or 0.0
        rel_max = _num(r[22]) or 0.0
        # search radius: seismic-only locations need a wide net; GT locations a small one
        radius_km = 2.0 if src == "ground_truth" else max(3.0, min(abs_err + rel_max, 25.0))

        qc = []
        if lon is not None and lon > 0:
            qc.append(f"positive longitude ({lon}) — missing minus sign?")
            lon = -abs(lon)
        if lon is not None and not (-180 <= lon <= -125):
            qc.append(f"longitude {lon} outside Alaska — typo?")
        if lat is not None and not (51 <= lat <= 72):
            qc.append(f"latitude {lat} outside Alaska — typo?")

        out.append(dict(
            event_id=str(r[0]).strip(),
            datetime_utc=when,
            lat=lat, lon=lon, loc_source=src,
            search_radius_km=round(radius_km, 1),
            has_ground_truth=str(r[9]).strip().upper().startswith("Y") if r[9] else False,
            in_esec=str(r[14]).strip() if r[14] else "",
            gt_volume_m3=_num(r[12]),
            seismic_volume_m3=_num(r[23]),
            coherency=_num(r[19]),
            abs_loc_err_km=_num(r[20]),
            nearest_station=str(r[15]).strip() if r[15] else "",
            qc_flags="; ".join(qc),
            notes=str(r[29]).strip() if r[29] else "",
        ))
    return pd.DataFrame(out)


def needs_groundtruth(df: pd.DataFrame) -> pd.DataFrame:
    """Events worth running imagery on: not confirmed in ESEC, or missing GT location/volume."""
    m = (~df.in_esec.str.upper().eq("Y")) | (~df.has_ground_truth) | df.gt_volume_m3.isna()
    return df[m & df.lat.notna() & df.datetime_utc.notna()].reset_index(drop=True)


if __name__ == "__main__":
    import sys
    df = load_events(sys.argv[1] if len(sys.argv) > 1 else "LandslideInventory.xlsx")
    df.to_csv("events.csv", index=False)
    print(df.to_string(max_colwidth=30))
    print(f"\n{len(df)} events parsed; {len(needs_groundtruth(df))} flagged for ground-truthing")
